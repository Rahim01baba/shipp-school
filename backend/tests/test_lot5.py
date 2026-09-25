"""Tests du lot 5 : contacts parents, vehicules et documents, navettes d'activite,
contrats incomplets, retenues, suivi mensuel Enko / SHIPP, tarifs, import ENKO.

Prerequis : base recreee avec les migrations 001 a 005.
Le fichier Excel de test est FICTIF (genere ici) : aucune donnee reelle d'eleve.
Usage : python3 test_lot5.py
"""
import datetime, io, json, sys, uuid, urllib.request, urllib.error
import openpyxl
from client import call, login, BASE, _tokens
from test_acces import sql, check, RESULTS

TODAY = datetime.date.today()
D = lambda n: (TODAY + datetime.timedelta(days=n)).isoformat()


def multipart(who, path, fields, filename, content):
    b = uuid.uuid4().hex
    parts = [f'--{b}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode() for k, v in fields.items()]
    parts.append(f'--{b}\r\nContent-Disposition: form-data; name="fichier"; filename="{filename}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode() + content + b'\r\n')
    parts.append(f'--{b}--\r\n'.encode())
    req = urllib.request.Request(BASE + path, data=b''.join(parts), method='POST')
    req.add_header('Content-Type', f'multipart/form-data; boundary={b}')
    req.add_header('Authorization', 'Bearer ' + login(who))
    try:
        with urllib.request.urlopen(req) as r:
            return r.status, json.loads(r.read().decode())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode() or '{}')


def fichier_enko_fictif():
    """Classeur fictif au format du suivi ENKO. Le nom de la 1re feuille est volontairement faux."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ANNEE 2023_2024'
    mois = ['Septembre 2026', 'octobre 2026', 'novembre 2026', 'décembre 2026', 'juillet 2027', 'août 2027']
    ws.append(['N°', 'Type de payment', 'Nom & Prénom', "Statut de l'inscription", 'Campus', 'Classe', 'Ligne', 'Commune', 'Localisation', 'Conducteurs', 'Vehicule', 'Montant à Payer'] + mois)
    ws.append([1, 'Trimestre', 'TESTA Alpha Beta', 'Validé / Début Sce', 'Enko Riviéra', 'MYP_5', 'ABIDJAN NORD', 'COCODY', 'Pres du rond-point', 'Richard', 'Van 10 pax', 72444.44, 'Enko, Shipp', 'Enko', 'Non', None, None, 'Enko'])
    ws.append([2, None, 'TESTB GAMMA', 'Validé / Début Sce', 'Enko Angre', 'PYP 3', 'BINGERVILLE', None, None, 'Richard', 'Van 10 pax', 80000, 'Shipp', 'Enko, Shipp', 'Arret S/c', None, None, None])
    ws.append([3, None, 'TESTC Delta', 'Validé / Début Sce, Service arrété', 'Enko Annexe', 'DP1', 'ABIDJAN NORD', None, None, 'Inconnu', 'Berline', 72444.44, 'Non', None, None, None, None, None])
    ws.append([4, None, 'TESTA Alpha Beta', 'Validé / Début Sce', 'Enko Riviera', 'MYP5', 'ABIDJAN NORD', None, None, None, None, 72444.44, None, None, None, None, None, None])
    ws.append([5, 'Annee', 'TESTD Epsilon', 'interréssé', 'Enko Riviera', 'GRADE 8', 'BINGERVILLE 1', None, None, None, None, None, None, None, None, None, None, None])
    ws2 = wb.create_sheet('Feuille sans eleves')
    ws2.append(['Jour', 'Nombre'])
    ws3 = wb.create_sheet('ANNEE 2026_2027')
    ws3.append(['Nom & Prénom', "Statut de l'inscription", 'Campus', 'Classe', 'Ligne', 'Montant à Payer', 'Septembre 2026'])
    ws3.append(['TESTE Zeta', 'Validé / Début Sce', 'Enko Annexe', 'DP2', 'BASSAM', 90000, 'Enko'])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    A = 'rahim'
    rh_id = call('POST', '/users.php', {'name': 'RH Test', 'email': 'rh@example.com', 'password': 'RhRecette2026!', 'role_key': 'rh'}, who=A)[1]['data']['id']
    sql(f'UPDATE users SET ecole_id = 2 WHERE id = {rh_id}')
    sql(f'UPDATE user_roles SET ecole_id = 2 WHERE user_id = {rh_id}')
    _tokens['rh'] = login('rh', password='RhRecette2026!', identifiant='rh@example.com')
    ch1 = int(sql('SELECT id FROM chauffeurs WHERE user_id = 4')[0]['id'])
    call('POST', '/vehicle-assignments.php', {'chauffeur_id': ch1, 'vehicle_id': 1, 'date_debut': D(-30)}, who='fleet')

    # ---------- Contacts parents ----------
    code, js = call('POST', '/eleve-contacts.php', {'eleve_id': 6, 'nom': 'Parent Fictif', 'lien': 'mere', 'telephone': '07 07 00 00 01', 'principal': True}, who=A)
    tel = sql(f"SELECT telephone FROM eleve_contacts WHERE id = {js.get('id', 0)}")
    check('L5-01', 'Contact parent ajoute, telephone normalise +225', code == 200 and tel and tel[0]['telephone'] == '+2250707000001', str((code, js, tel)))
    code, js = call('POST', '/eleve-contacts.php', {'eleve_id': 6, 'nom': 'Sans moyen'}, who=A)
    check('L5-02', 'Refus : contact sans telephone ni e-mail', code == 400)
    code, js = call('GET', '/eleve-contacts.php?eleve_id=6', who='parent')
    code2, js2 = call('GET', '/eleve-contacts.php?eleve_id=3', who='parent')
    check('L5-03', "Parent : voit les contacts de son enfant, pas ceux d'un autre", code == 200 and len(js['data']) == 1 and code2 == 200 and js2['data'] == [], str((js, js2)))
    code, _ = call('POST', '/eleve-contacts.php', {'eleve_id': 6, 'telephone': '0101010101'}, who='fleet')
    code2, _ = call('GET', '/eleve-contacts.php?eleve_id=6', who='chauffeur')
    check('L5-04', 'Fleet en lecture seule ; chauffeur sans acces aux contacts', code == 403 and code2 == 403, str((code, code2)))

    # ---------- Vehicule : identification et documents ----------
    code, _ = call('PUT', '/crud.php?module=vehicules', {'id': 1, 'marque': 'Toyota', 'annee': 2018, 'type_vehicule': 'Van 10 pax', 'proprietaire': 'SHIPP'}, who=A)
    v = sql('SELECT marque, annee, type_vehicule FROM vehicules WHERE id = 1')[0]
    check('L5-05', 'Vehicule complete (marque, annee, type)', code == 200 and v['marque'] == 'Toyota' and int(v['annee']) == 2018, str(v))
    code, js = call('POST', '/vehicle-documents.php', {'vehicle_id': 1, 'type': 'assurance', 'numero': 'ASS-1', 'organisme': 'Assureur fictif', 'date_expiration': D(10)}, who='fleet')
    doc = js.get('id')
    code2, _ = call('PUT', '/vehicle-documents.php', {'id': doc, 'action': 'valider'}, who='fleet')
    check('L5-06', 'Fleet enregistre et valide une assurance', code == 200 and code2 == 200, str(js))
    al = call('GET', '/alertes.php', who='fleet')[1]['data']
    types = {(a['type'], a.get('entite_id')) for a in al}
    check('L5-07', 'Alertes vehicule : assurance bientot expiree, visite technique manquante',
          ('document_vehicule_expire_bientot', 1) in types and ('document_vehicule_manquant', 1) in types, str([a for a in al if 'vehicule' in a['type']])[:400])
    code, js = call('GET', '/vehicle-documents.php?vehicle_id=1', who='chauffeur')
    code2, js2 = call('GET', '/vehicle-documents.php?vehicle_id=2', who='chauffeur')
    check('L5-08', 'Chauffeur : documents de son vehicule seulement', code == 200 and len(js['data']) == 1 and js2.get('data') == [], str((js, js2))[:300])
    code, _ = call('POST', '/vehicle-documents.php', {'vehicle_id': 1, 'type': 'assurance'}, who='chauffeur')
    check('L5-09', 'Chauffeur : pas de depot de document vehicule', code == 403)

    # ---------- Navettes d'activite ----------
    dom = call('POST', '/crud.php?module=circuits', {'nom': 'DOMICILE-1', 'statut': 'actif'}, who=A)[1]['id']
    act = call('POST', '/crud.php?module=circuits', {'nom': 'NATATION', 'statut': 'actif', 'type_circuit': 'activite', 'activite': 'Natation',
                                                     'destination': 'Club fictif', 'jours_semaine': '3', 'heure_depart': '14:30', 'heure_retour': '15:50'}, who=A)[1]['id']
    for c in (dom, act):
        call('POST', '/crud.php?module=etapes', {'circuit_id': c, 'nom': 'Ecole', 'ordre': 1, 'heure_estimee': '07:00:00', 'statut': 'active'}, who=A)
    call('POST', '/eleve-affectations.php', {'eleve_id': 6, 'circuit_id': dom}, who='fleet')
    code, _ = call('POST', '/eleve-affectations.php', {'eleve_id': 6, 'circuit_id': act}, who='fleet')
    actives = sql("SELECT circuit_id FROM eleve_affectations_transport WHERE eleve_id = 6 AND statut = 'active'")
    ecid = sql('SELECT circuit_id FROM eleves WHERE id = 6')[0]['circuit_id']
    check('L5-10', "Navette d'activite ajoutee sans remplacer le circuit domicile", code == 200 and {int(a['circuit_id']) for a in actives} == {dom, act} and int(ecid) == dom, str((actives, ecid)))
    mer = TODAY + datetime.timedelta(days=(2 - TODAY.weekday()) % 7 or 7)
    jeu = mer + datetime.timedelta(days=1)
    js_m = call('POST', '/trajets-generer.php', {'date': mer.isoformat(), 'sens': 'aller', 'circuit_ids': [dom, act]}, who=A)[1]
    js_j = call('POST', '/trajets-generer.php', {'date': jeu.isoformat(), 'sens': 'aller', 'circuit_ids': [dom, act]}, who=A)[1]
    check('L5-11', 'Navette generee seulement le mercredi', len(js_m.get('crees', [])) == 2 and len(js_j.get('crees', [])) == 1, str((js_m, js_j))[:300])
    code, pe = call('GET', '/parent-enfants.php', who='parent')
    enfant = [e for e in pe.get('data', []) if int(e['eleve']['id']) == 6]
    check('L5-12', 'Espace parent : circuit domicile + navette listee', code == 200 and enfant and enfant[0]['affectation'] and len(enfant[0]['activites']) == 1, str(pe)[:300])

    # ---------- Contrat : champs remplis au besoin ----------
    call('POST', '/chauffeur-documents.php', {'chauffeur_id': ch1, 'type': 'cni', 'numero': 'CNI-FICTIVE-1'}, who='rh')
    tpl = call('POST', '/contract-templates.php', {'code': 'CUV', 'titre': 'Contrat', 'contenu': 'CNI {{chauffeur_cni}} ; vehicule {{vehicule_marque}} {{vehicule_modele}} {{vehicule_annee}} {{vehicule_immatriculation}} ; duree {{duree_mois}}'}, who='rh')[1]['id']
    call('PUT', '/contract-templates.php', {'id': tpl, 'action': 'activer'}, who='rh')
    code, js = call('POST', '/chauffeur-contracts.php', {'chauffeur_id': ch1, 'vehicle_id': 1, 'template_id': tpl, 'remuneration_montant': 150000}, who='rh')
    ct = js.get('id')
    check('L5-13', 'Contrat brouillon cree sans date ni duree (champs vides)', code == 200 and 'duree_mois' in js.get('variables_manquantes', []), str(js))
    contenu = sql(f'SELECT contenu_genere FROM chauffeur_contracts WHERE id = {ct}')[0]['contenu_genere']
    check('L5-14', 'Variables CNI, marque et annee du vehicule remplies', 'CNI-FICTIVE-1' in contenu and 'Toyota' in contenu and '2018' in contenu, contenu)
    code, js = call('PUT', '/chauffeur-contracts.php', {'id': ct, 'action': 'activer'}, who='rh')
    check('L5-15', "Activation refusee tant que la date de debut n'est pas renseignee", code == 409, str(js))
    call('PUT', '/chauffeur-contracts.php', {'id': ct, 'action': 'modifier', 'date_debut': D(-60)}, who='rh')
    code, _ = call('PUT', '/chauffeur-contracts.php', {'id': ct, 'action': 'activer'}, who='rh')
    check('L5-16', 'Activation apres saisie de la date', code == 200)

    # ---------- Retenues chiffrees ----------
    periode = TODAY.strftime('%Y-%m')
    code, js = call('POST', '/chauffeur-retenues.php', {'chauffeur_id': ch1, 'periode': periode, 'montant': 20000, 'motif': 'Amende non reglee'}, who='rh')
    ret = js.get('id')
    check('L5-17', 'RH propose une retenue chiffree', code == 200 and ret, str(js))
    inc = call('POST', '/incidents.php', {'titre': 'Rayure', 'categorie': 'accident', 'chauffeur_id': ch1, 'vehicule_id': 1}, who='fleet')[1]['id']
    code, js = call('POST', '/chauffeur-retenues.php', {'chauffeur_id': ch1, 'periode': periode, 'montant': 5000, 'motif': 'Rayure', 'incident_id': inc}, who='rh')
    check('L5-18', 'Refus : retenue sur un accident non qualifie « chauffeur »', code == 409, str(js))
    call('PUT', '/incidents.php', {'id': inc, 'action': 'qualifier', 'responsabilite': 'chauffeur', 'commentaire': 'Manoeuvre en marche arriere'}, who='fleet')
    code, _ = call('POST', '/chauffeur-retenues.php', {'chauffeur_id': ch1, 'periode': periode, 'montant': 5000, 'motif': 'Rayure', 'incident_id': inc}, who='rh')
    check('L5-19', 'Retenue acceptee apres qualification humaine', code == 200)
    code, _ = call('PUT', '/chauffeur-retenues.php', {'id': ret, 'action': 'valider'}, who='rh')
    code2, etat = call('GET', f'/chauffeur-retenues.php?etat={periode}', who='rh')
    row = [r for r in etat.get('data', []) if int(r['chauffeur_id']) == ch1]
    check('L5-20', 'Etat mensuel : 150 000 - 20 000 validee = 130 000 (retenue proposee non deduite)',
          code == 200 and row and row[0]['net'] == 130000 and int(row[0]['en_attente']) == 1, str(row))
    c1, _ = call('GET', f'/chauffeur-retenues.php?chauffeur_id={ch1}', who='chauffeur')
    c2, _ = call('GET', f'/chauffeur-retenues.php?chauffeur_id={ch1}', who='fleet')
    check('L5-21', 'Retenues invisibles pour le chauffeur et la flotte', c1 == 403 and c2 == 403, str((c1, c2)))

    # ---------- Tarifs et suivi mensuel Enko / SHIPP ----------
    code, _ = call('POST', '/tarifs.php', {'annee_scolaire_id': 1, 'zone': 'ABIDJAN NORD', 'montant_mensuel': 72444.44}, who=A)
    code2, _ = call('POST', '/tarifs.php', {'annee_scolaire_id': 1, 'zone': 'ABIDJAN NORD', 'montant_mensuel': 1}, who=A)
    check('L5-22', 'Tarif par zone et par annee, sans doublon', code == 200 and code2 == 409)
    call('PUT', '/crud.php?module=abonnements', {'id': 4, 'zone_tarifaire': 'ABIDJAN NORD'}, who=A)
    code, js = call('POST', '/echeances-transport.php', {'action': 'generer'}, who=A)
    n6 = sql('SELECT COUNT(*) AS n, MIN(montant) AS m FROM echeances_transport WHERE eleve_id = 6')[0]
    check('L5-23', "Mois generes sur l'annee scolaire (sept. -> juil.), montant du tarif de zone", code == 200 and int(n6['n']) == 11 and float(n6['m']) == 72444.44, str((js, n6)))
    code, _ = call('PUT', '/echeances-transport.php', {'eleve_id': 6, 'mois': '2026-10', 'encaisse_enko': True}, who=A)
    g = call('GET', '/echeances-transport.php', who=A)[1]
    check('L5-24', "Mois encaisse par Enko et non recu par SHIPP : compte « a reverser »", code == 200 and g['totaux']['2026-10']['a_reverser'] == 72444.44, str(g['totaux'].get('2026-10')))
    call('PUT', '/echeances-transport.php', {'eleve_id': 6, 'mois': '2026-10', 'recu_shipp': True}, who=A)
    g = call('GET', '/echeances-transport.php', who=A)[1]
    h = sql("SELECT champ, nouvelle_valeur FROM echeances_historique ORDER BY id")
    check('L5-25', 'Reception SHIPP enregistree, historique champ par champ', g['totaux']['2026-10']['a_reverser'] == 0 and g['totaux']['2026-10']['shipp'] == 72444.44
          and [x['champ'] for x in h] == ['encaisse_enko', 'recu_shipp'], str(h))
    code, js = call('GET', '/echeances-transport.php', who='parent')
    code2, _ = call('PUT', '/echeances-transport.php', {'eleve_id': 6, 'mois': '2026-11', 'recu_shipp': True}, who='parent')
    check('L5-26', 'Parent : voit les mois de son enfant uniquement, sans modification', code == 200 and [int(e['id']) for e in js['data']] == [6] and js['droits']['modifier'] is False and code2 == 403, str((code2, js.get('droits'))))
    sql("INSERT INTO annees_scolaires (id, ecole_id, libelle, date_debut, date_fin, statut) VALUES (2, 2, '2025-2026', '2025-09-01', '2026-06-30', 'archivee')")
    sql("INSERT INTO echeances_transport (ecole_id, annee_scolaire_id, eleve_id, mois, montant, encaisse_enko) VALUES (2, 2, 3, '2026-03-01', 80000, 1)")
    al = call('GET', '/alertes.php', who=A)[1]['data']
    check('L5-27', "Alerte : mois passe encaisse par Enko, non recu par SHIPP", any(a['type'] == 'paiements_a_reverser' for a in al), str([a['type'] for a in al]))
    code, js = call('GET', '/reporting.php?rapport=paiements&date_debut=2026-03-01&date_fin=2026-10-31', who=A)
    mars = [r for r in js.get('data', []) if r['mois'] == '2026-03']
    code2, js2 = call('GET', '/reporting.php?liste=a_reverser&date_debut=2026-03-01&date_fin=2026-10-31', who=A)
    check('L5-28', 'Reporting paiements par mois + detail « a reverser »', code == 200 and mars and mars[0]['a_reverser'] == 80000 and code2 == 200 and len(js2['data']) == 1, str((js.get('data'), js2.get('data')))[:300])

    # ---------- Import ENKO (fichier fictif) ----------
    contenu = fichier_enko_fictif()
    code, _ = multipart('fleet', '/imports.php', {}, 'suivi.xlsx', contenu)
    check('L5-29', 'Import reserve a l\'administration', code == 403)
    code, js = multipart(A, '/imports.php', {}, 'suivi.xlsx', contenu)
    lot = js.get('lot_id')
    f1 = [f for f in js.get('feuilles', []) if f['nom'] == 'ANNEE 2023_2024']
    check('L5-30', "Analyse : annee suggeree d'apres les mois (2026-2027), pas d'apres le nom de feuille",
          code == 200 and f1 and f1[0]['lignes'] == 5 and f1[0]['suggestion_annees'][0]['libelle'] == '2026-2027', str(js)[:400])
    sans = [f for f in js['feuilles'] if f['nom'] == 'Feuille sans eleves']
    check('L5-31', 'Feuille sans colonne eleve signalee non importable', sans and sans[0]['importable'] is False)
    code, js = call('PUT', '/imports.php', {'lot_id': lot, 'action': 'preparer', 'feuille': 'ANNEE 2023_2024'}, who=A)
    check('L5-32', "Preparation refusee sans annee choisie", code == 400, str(js))
    code, js = call('PUT', '/imports.php', {'lot_id': lot, 'action': 'preparer', 'feuille': 'ANNEE 2023_2024', 'annee_scolaire_id': 1, 'options': {'circuits_par_conducteur': True}}, who=A)
    st = js.get('stats', {})
    check('L5-33', 'Classement : 1 doublon en erreur, conducteurs sans correspondance et statut ambigu = donnee insuffisante',
          code == 200 and st.get('erreur') == 1 and st.get('insuffisant') == 3 and st.get('mois_hors_annee', 0) >= 1, str(st)[:500])
    lignes = call('GET', f'/imports.php?lot_id={lot}', who=A)[1]['lignes']
    l1 = [l for l in lignes if l['donnees']['source_nom'] == 'TESTA Alpha Beta'][0]
    check('L5-34', 'Normalisation : nom/prenom, classe MYP5, campus, mois Enko/Shipp, mois hors annee ecarte',
          l1['donnees']['nom'] == 'TESTA' and l1['donnees']['prenom'] == 'Alpha Beta' and l1['donnees']['classe'] == 'MYP5' and l1['donnees']['campus'] == 'Enko Riviera'
          and l1['donnees']['mois'].get('2026-09') == {'enko': 1, 'shipp': 1, 'arret': 0} and '2027-08' not in l1['donnees']['mois'], json.dumps(l1['donnees'])[:400])
    call('PUT', '/imports.php', {'action': 'correspondances', 'ecole_id': 2, 'items': [{'type': 'conducteur', 'valeur_source': 'Richard', 'cible_id': ch1}]}, who=A)
    code, js = call('PUT', '/imports.php', {'lot_id': lot, 'action': 'preparer', 'feuille': 'ANNEE 2023_2024', 'annee_scolaire_id': 1, 'options': {'circuits_par_conducteur': True}}, who=A)
    check('L5-35', 'Apres correspondance « Richard -> chauffeur », seules les lignes ambigues restent insuffisantes', js.get('stats', {}).get('insuffisant') == 1, str(js.get('stats'))[:300])
    code, js = call('PUT', '/imports.php', {'lot_id': lot, 'action': 'valider'}, who=A)
    cr = js.get('crees', {})
    check('L5-36', 'Validation : 4 eleves, 4 abonnements, 1 circuit par conducteur, 2 affectations, tarifs de zone',
          code == 200 and cr.get('eleves') == 4 and cr.get('abonnements') == 4 and cr.get('circuits') == 1 and cr.get('affectations') == 2 and cr.get('tarifs') == 1, str(js))
    e = sql(f"SELECT e.id, e.annee_scolaire_id, a.statut, a.zone_tarifaire, a.periodicite FROM eleves e JOIN abonnements a ON a.eleve_id = e.id AND a.type = 'transport' WHERE e.import_lot_id = {lot} ORDER BY e.id")
    check('L5-37', 'Eleves rattaches a l\'annee choisie ; statut « arrete / valide » ambigu -> en attente ; prospect -> en attente',
          all(int(x['annee_scolaire_id']) == 1 for x in e) and [x['statut'] for x in e] == ['actif', 'actif', 'en_attente', 'en_attente'] and e[0]['periodicite'] == 'trimestrielle', str(e))
    ech = sql(f"SELECT DATE_FORMAT(mois, '%Y-%m') AS m, encaisse_enko, recu_shipp, arret_service FROM echeances_transport WHERE import_lot_id = {lot} AND eleve_id = {e[1]['id']} ORDER BY mois")
    check('L5-38', 'Mois importes : « Shipp », « Enko, Shipp », « Arret S/c »',
          [(x['m'], int(x['encaisse_enko']), int(x['recu_shipp']), int(x['arret_service'])) for x in ech] == [('2026-09', 0, 1, 0), ('2026-10', 1, 1, 0), ('2026-11', 0, 0, 1)], str(ech))
    code, js = multipart(A, '/imports.php', {}, 'suivi-copie.xlsx', contenu)
    lot2 = js.get('lot_id')
    f1 = [f for f in js['feuilles'] if f['nom'] == 'ANNEE 2023_2024'][0]
    code2, js2 = call('PUT', '/imports.php', {'lot_id': lot2, 'action': 'preparer', 'feuille': 'ANNEE 2023_2024', 'annee_scolaire_id': 1}, who=A)
    check('L5-39', 'Meme fichier, meme feuille : import en double refuse', f1['deja_importee'] is True and code2 == 409, str(js2))
    code, js = call('PUT', '/imports.php', {'lot_id': lot2, 'action': 'preparer', 'feuille': 'ANNEE 2026_2027', 'annee_scolaire_id': 1}, who=A)
    call('PUT', '/imports.php', {'lot_id': lot2, 'action': 'valider'}, who=A)
    code, js = call('PUT', '/imports.php', {'lot_id': lot2, 'action': 'annuler', 'motif': 'Essai'}, who=A)
    reste = sql(f'SELECT COUNT(*) AS n FROM eleves WHERE import_lot_id = {lot2}')[0]['n']
    check('L5-40', "Annulation d'un lot sans dependance : ses eleves sont retires", code == 200 and int(reste) == 0, str(js))
    call('PUT', '/echeances-transport.php', {'eleve_id': e[0]['id'], 'mois': '2026-12', 'recu_shipp': True}, who=A)
    code, js = call('PUT', '/imports.php', {'lot_id': lot, 'action': 'annuler', 'motif': 'Essai'}, who=A)
    check('L5-41', "Annulation refusee des qu'une donnee s'est rattachee au lot (mois modifie)", code == 409 and 'mois saisis' in js.get('message', ''), str(js))
    j = sql("SELECT module_key, COUNT(*) AS n FROM journal_activite WHERE module_key IN ('eleve_contacts','vehicle_documents','chauffeur_retenues','echeances_transport','tarifs','imports') GROUP BY module_key")
    check('L5-42', 'Actions du lot 5 journalisees', len(j) == 6, str(j))

    ok = sum(1 for r in RESULTS if r[2])
    print(f'\n{ok}/{len(RESULTS)} tests reussis')
    sys.exit(0 if ok == len(RESULTS) else 1)


if __name__ == '__main__':
    main()
